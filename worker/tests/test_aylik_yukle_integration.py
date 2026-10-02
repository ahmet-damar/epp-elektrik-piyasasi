"""EPP — `worker/scripts/aylik_yukle.py` entegrasyon testi
(`Claude outputs/PROMPT_AYLIK_YUKLEME_SCRIPTI_2026-10-02.md`).

`worker/tests/test_parser.py:_sentetik_workbook()`'u (gerçek dosyaya karşı
doğrulanmış) yeniden kullanır — sentinel tarih_id (`999801`, yıl 9998,
established desen) ile gerçek veriyle ÇAKIŞMAZ. Dosya-keşif testi için bir
başlık sayfası eklenir; "negatif satır sayısı" testleri için T11'in
ESKİŞEHİR satırındaki hücreler kontrollü şekilde negatife çevrilir (her
negatif hücre tam olarak 1 red satırına karşılık gelir — ayrı il eklemeye
gerek yok).
"""

from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

import openpyxl
import psycopg
import pytest

from worker.scripts import aylik_yukle
from worker.tests.test_parser import _sentetik_workbook

DATABASE_URL = os.environ.get("DATABASE_URL")

pytestmark = pytest.mark.skipif(
    not DATABASE_URL,
    reason="DATABASE_URL tanımlı değil (yalnız CI 'integration' job'ında çalışır)",
)

_TEST_TARIH_ID = 999801  # yıl 9998 — established sentinel deseni
_TEST_YIL = 9998
_TEST_AY_ADI = "Ocak"
_TEST_SOURCE_PERIOD = "9998-01"


@pytest.fixture
def conn():  # type: ignore[no-untyped-def]
    assert DATABASE_URL is not None
    with psycopg.connect(DATABASE_URL) as connection:
        yield connection
        connection.rollback()
        # Bazı testler (--uygula) GERÇEKTEN commit eder (script'in kendisi
        # öyle çalışıyor) — sentinel tarih_id'ye bağlı HER ŞEY burada elle
        # silinir (bkz. test_job_worker_integration.py'nin AYNI deseni).
        with connection.cursor() as cur:
            for tablo in (
                "fact_tuketim",
                "fact_abone",
                "fact_serbest_tuketici",
                "fact_uretim",
                "fact_tuketim_ulke_geneli",
                "fact_uretim_kaynak_geneli",
                "fact_uretim_il_geneli",
            ):
                cur.execute(
                    f"DELETE FROM {tablo} WHERE tarih_id = %s",  # nosec B608 - sabit tablo listesi
                    (_TEST_TARIH_ID,),
                )
            cur.execute(
                """
                DELETE FROM audit_log WHERE record_id IN (
                    SELECT ib.batch_id FROM ingestion_batch ib
                    JOIN source_asset sa ON sa.source_asset_id = ib.source_asset_id
                    WHERE sa.source_period = %s
                )
                """,
                (_TEST_SOURCE_PERIOD,),
            )
            cur.execute(
                """
                DELETE FROM ingestion_batch WHERE source_asset_id IN (
                    SELECT source_asset_id FROM source_asset WHERE source_period = %s
                )
                """,
                (_TEST_SOURCE_PERIOD,),
            )
            cur.execute(
                "DELETE FROM source_asset WHERE source_period = %s",
                (_TEST_SOURCE_PERIOD,),
            )
            cur.execute("DELETE FROM dim_tarih WHERE tarih_id = %s", (_TEST_TARIH_ID,))
        connection.commit()


def _workbook(negatif_hucre_sayisi: int = 0) -> openpyxl.Workbook:
    """`_sentetik_workbook()`'a bir başlık sayfası ekler (dosya-keşfi için)
    ve T11'in ESKİŞEHİR satırında `negatif_hucre_sayisi` kadar hücreyi
    (Tarımsal'dan başlayarak) negatife çevirir — her negatif hücre, P0
    kuralıyla (negatif değer → reddet) tam 1 `fact_tuketim` red satırı
    üretir (il bazında kontrollü bir sayı, yeni il eklemeye gerek yok)."""
    wb = _sentetik_workbook()
    ic = wb.create_sheet("İçindekiler", 0)
    ic.append(
        [
            None,
            f"ELEKTRİK PİYASASI SEKTÖR RAPORU {_TEST_AY_ADI.upper()} {_TEST_YIL} - EK",
        ]
    )

    # _sentetik_workbook() KASITLI OLARAK 1 negatif satır taşıyor (T13,
    # "ST Olma Hakkını Kullanmayan Aboneler" satırı, Mesken=-5.0,
    # Toplam=-5.0 — dogrula_serbest_tuketici'nin kendi red testi için).
    # Bu fonksiyonun "temiz" (0 negatif) tabanı için önce bu nötrlenir;
    # negatif_hucre_sayisi SADECE T11'e (fact_tuketim) kontrollü olarak
    # eklenen EK negatif hücre sayısıdır.
    t13 = wb["Tablo 13"]
    t13.cell(row=6, column=3).value = 5.0
    t13.cell(row=6, column=8).value = 5.0

    if negatif_hucre_sayisi:
        t11 = wb["Tablo 11"]
        # satır 3 = ESKİŞEHİR; sütunlar: 1=İL,2=Aydınlatma,3=Kamu,4=Mesken,
        # 5=Sanayi-DAĞITIM,6=Sanayi-İLETİM,7=Tarımsal — sondan başlanır.
        hedef_kolonlar = [7, 6, 5, 4, 2][:negatif_hucre_sayisi]
        for kolon in hedef_kolonlar:
            mevcut = t11.cell(row=3, column=kolon).value
            deger = float(str(mevcut).replace(".", "").replace(",", "."))
            t11.cell(row=3, column=kolon).value = f"-{deger:.2f}".replace(".", ",")
    return wb


def _wb_bytes(wb: openpyxl.Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _dosya_yaz(tmp_path: Path, negatif_hucre_sayisi: int = 0) -> Path:
    yol = tmp_path / "test_aylik_yukle.xlsx"
    yol.write_bytes(_wb_bytes(_workbook(negatif_hucre_sayisi)))
    return yol


# ---------------------------------------------------------------------------
# kaynak_dosya_bul() — DB gerektirmez, saf dosya-keşif mantığı
# ---------------------------------------------------------------------------


def test_kaynak_dosya_bul_tek_eslesme_bulur(tmp_path: Path) -> None:
    yol = _dosya_yaz(tmp_path)
    bulunan = aylik_yukle.kaynak_dosya_bul(tmp_path, _TEST_YIL, _TEST_AY_ADI)
    assert bulunan == yol


def test_kaynak_dosya_bul_sifir_eslesme_durur(tmp_path: Path) -> None:
    (tmp_path / "ilgisiz.xlsx").write_bytes(_wb_bytes(openpyxl.Workbook()))
    with pytest.raises(aylik_yukle.DurdurmaHatasi, match="bulunamadı"):
        aylik_yukle.kaynak_dosya_bul(tmp_path, _TEST_YIL, _TEST_AY_ADI)


def test_kaynak_dosya_bul_birden_fazla_eslesme_durur(tmp_path: Path) -> None:
    (tmp_path / "a.xlsx").write_bytes(_wb_bytes(_workbook()))
    (tmp_path / "b.xlsx").write_bytes(_wb_bytes(_workbook()))
    with pytest.raises(aylik_yukle.DurdurmaHatasi, match="BİRDEN FAZLA"):
        aylik_yukle.kaynak_dosya_bul(tmp_path, _TEST_YIL, _TEST_AY_ADI)


# ---------------------------------------------------------------------------
# calistir() — tam zincir, gerçek DB'ye karşı
# ---------------------------------------------------------------------------


def _aktif_satir_sayisi(conn, tablo: str, tarih_id: int) -> int:  # type: ignore[no-untyped-def]
    with conn.cursor() as cur:
        cur.execute(
            f"select count(*) from {tablo} where is_active and tarih_id = %s",  # nosec B608
            (tarih_id,),
        )
        row = cur.fetchone()
        assert row is not None
        return row[0]  # type: ignore[no-any-return]


def test_dry_run_hicbir_sey_yazmaz(conn, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    """--uygula verilmeden çağrıldığında (uygula=False): zincirler GERÇEKTEN
    işlenir (aynı sadakatle) ama sonunda rollback edilir — hiçbir satır
    kalıcı olmaz, 'hedef ay dışı fark' listesi de (zaten hiç kalıcı veri
    yok) boş olmalı."""
    yol = _dosya_yaz(tmp_path)

    sonuclar, farklar = aylik_yukle.calistir(
        conn,
        yol=yol,
        tarih_id=_TEST_TARIH_ID,
        source_period=_TEST_SOURCE_PERIOD,
        parser_version="aylik-yukle-test-v1",
        uygula=False,
        negatif_red_esigi=0,
        elle_onay=None,
        actor="test-suite",
    )

    assert farklar == []
    assert sonuclar[0]["zincir"] == "ana"
    conn.commit()  # fixture'ın kendi rollback'i zaten temiz ama açıkça kapatalım
    assert _aktif_satir_sayisi(conn, "fact_tuketim", _TEST_TARIH_ID) == 0
    assert _aktif_satir_sayisi(conn, "fact_tuketim_ulke_geneli", _TEST_TARIH_ID) == 0
    assert _aktif_satir_sayisi(conn, "fact_uretim_kaynak_geneli", _TEST_TARIH_ID) == 0


def test_uygula_temiz_veriyle_tum_zinciri_otomatik_aktive_eder(
    conn,
    tmp_path: Path,  # type: ignore[no-untyped-def]
) -> None:
    """Sıfır red/karantina ile (negatif_hucre_sayisi=0) her üç zincir de
    `otomatik_onaya_uygun()`/`periyot_aktivasyona_uygun_mu()`'dan True almalı
    ve ELLE ONAY GEREKMEDEN aktive olmalı — 'temiz ay sorunsuz akar' kuralı."""
    yol = _dosya_yaz(tmp_path, negatif_hucre_sayisi=0)

    sonuclar, farklar = aylik_yukle.calistir(
        conn,
        yol=yol,
        tarih_id=_TEST_TARIH_ID,
        source_period=_TEST_SOURCE_PERIOD,
        parser_version="aylik-yukle-test-v1",
        uygula=True,
        negatif_red_esigi=0,
        elle_onay=None,
        actor="test-suite",
    )

    assert farklar == []
    durumlar = {r["zincir"]: r["durum"] for r in sonuclar}
    assert durumlar["ana"] == "aktive_edildi"
    assert durumlar["uretim_geneli"] == "aktive_edildi"
    assert durumlar["ulke_geneli"] == "aktive_edildi"
    assert _aktif_satir_sayisi(conn, "fact_tuketim", _TEST_TARIH_ID) > 0
    assert _aktif_satir_sayisi(conn, "fact_tuketim_ulke_geneli", _TEST_TARIH_ID) > 0
    assert _aktif_satir_sayisi(conn, "fact_uretim_kaynak_geneli", _TEST_TARIH_ID) > 0


def test_negatif_red_esigi_asilinca_durur_kasitli_bozuk_senaryo(
    conn,
    tmp_path: Path,  # type: ignore[no-untyped-def]
) -> None:
    """KASITLI BOZUK ÖRNEK (yeni kabul kuralı): 2 negatif hücre enjekte
    edilip eşik 1'de tutulursa, ana zincir GERÇEKTEN durmalı — hiçbir
    tablo aktive edilmemeli, batch 'onay_bekliyor'da kalmalı."""
    yol = _dosya_yaz(tmp_path, negatif_hucre_sayisi=2)

    sonuclar, farklar = aylik_yukle.calistir(
        conn,
        yol=yol,
        tarih_id=_TEST_TARIH_ID,
        source_period=_TEST_SOURCE_PERIOD,
        parser_version="aylik-yukle-test-v1",
        uygula=True,
        negatif_red_esigi=1,
        elle_onay=None,
        actor="test-suite",
    )

    assert farklar == []
    assert sonuclar[0]["zincir"] == "ana"
    assert sonuclar[0]["durum"] == "DURDU_esik_asimi"
    assert sonuclar[0]["toplam_red"] == 2
    assert len(sonuclar) == 1  # üretim/ülke-geneli zincirlerine HİÇ geçilmedi

    with conn.cursor() as cur:
        cur.execute(
            "select status, error_summary from ingestion_batch where batch_id = %s",
            (sonuclar[0]["batch_id"],),
        )
        row = cur.fetchone()
        assert row is not None
        assert row[0] == "onay_bekliyor"
        assert "eşiği aşıldı" in row[1]
    assert _aktif_satir_sayisi(conn, "fact_tuketim", _TEST_TARIH_ID) == 0


def test_otomatik_onay_tutmazsa_gerekcesiz_elle_onay_mumkun_degil(
    conn,
    tmp_path: Path,  # type: ignore[no-untyped-def]
) -> None:
    """1 negatif hücre (eşik=1 ile GEÇİLİR, ama bu da red>0 ürettiğinden
    `otomatik_onaya_uygun()` yine de False döner) + `--elle-onay` YOK ->
    DURMALI, batch 'onay_bekliyor' kalmalı, HİÇ aktive edilmemeli."""
    yol = _dosya_yaz(tmp_path, negatif_hucre_sayisi=1)

    sonuclar, farklar = aylik_yukle.calistir(
        conn,
        yol=yol,
        tarih_id=_TEST_TARIH_ID,
        source_period=_TEST_SOURCE_PERIOD,
        parser_version="aylik-yukle-test-v1",
        uygula=True,
        negatif_red_esigi=1,
        elle_onay=None,
        actor="test-suite",
    )

    assert farklar == []
    assert sonuclar[0]["durum"] == "DURDU_elle_onay_gerekli"
    assert _aktif_satir_sayisi(conn, "fact_tuketim", _TEST_TARIH_ID) == 0


def test_gerekceli_elle_onay_aktive_eder_ve_audit_loga_yazar(
    conn,
    tmp_path: Path,  # type: ignore[no-untyped-def]
) -> None:
    """AYNI senaryo (1 negatif hücre, otomatik onay tutmuyor) ama
    `--elle-onay "<gerekçe>"` verilince ana zincir aktive OLMALI ve
    gerekçe `audit_log`'a yazılmalı."""
    yol = _dosya_yaz(tmp_path, negatif_hucre_sayisi=1)

    sonuclar, farklar = aylik_yukle.calistir(
        conn,
        yol=yol,
        tarih_id=_TEST_TARIH_ID,
        source_period=_TEST_SOURCE_PERIOD,
        parser_version="aylik-yukle-test-v1",
        uygula=True,
        negatif_red_esigi=1,
        elle_onay="test gerekçesi: bilinen desen",
        actor="test-suite",
    )

    assert farklar == []
    assert sonuclar[0]["durum"] == "aktive_edildi_elle_onay"
    assert _aktif_satir_sayisi(conn, "fact_tuketim", _TEST_TARIH_ID) > 0

    with conn.cursor() as cur:
        cur.execute(
            """
            SELECT payload FROM audit_log
            WHERE table_name = 'ingestion_batch' AND record_id = %s
              AND payload->>'olay' = 'elle_onay'
            """,
            (sonuclar[0]["batch_id"],),
        )
        row = cur.fetchone()
        assert row is not None
        assert row[0]["gerekce"] == "test gerekçesi: bilinen desen"


def test_main_bos_gerekceli_elle_onay_reddedilir(monkeypatch, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    """`--elle-onay ""` (boş string) argparse seviyesinde reddedilmeli —
    script'e hiç ulaşmadan, DB'ye dokunmadan."""
    monkeypatch.setattr(
        "sys.argv",
        ["aylik_yukle.py", "--ay", "999801", "--elle-onay", "   "],
    )
    assert aylik_yukle.main() == 1


def test_main_bos_gerekceli_yedek_atla_reddedilir(monkeypatch, tmp_path: Path) -> None:  # type: ignore[no-untyped-def]
    monkeypatch.setattr(
        "sys.argv",
        ["aylik_yukle.py", "--ay", "999801", "--yedek-atla", ""],
    )
    assert aylik_yukle.main() == 1
